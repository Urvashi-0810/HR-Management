"""
routes.py — FastAPI routes for the HR Management System backend.

Exposes RESTful endpoints for each pipeline step, designed for
the React frontend to consume.

API Groups:
  /api/email       — Email download & status
  /api/extract     — CV data extraction
  /api/candidates  — Read & shortlist candidates
  /api/job-match   — Job fit scoring
  /api/config      — Runtime configuration
"""

import logging
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel, Field

from src.config import (
    CV_FOLDER,
    OUTPUT_EXCEL,
    SHORTLIST_EXCEL,
    JOB_MATCH_EXCEL,
    LLM_PROVIDER,
    SUPPORTED_EXTENSIONS,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api")


# ──────────────────────────────────────────────
# Pydantic models (request/response schemas)
# ──────────────────────────────────────────────

class EmailDownloadRequest(BaseModel):
    """Optional overrides for email download."""
    imap_server: Optional[str] = None
    imap_port: Optional[int] = None
    email_address: Optional[str] = None
    email_password: Optional[str] = None
    email_folder: Optional[str] = None


class ExtractRequest(BaseModel):
    """Options for CV extraction."""
    provider: Optional[str] = Field(None, description="LLM provider: 'claude' or 'gemini'")
    folder: Optional[str] = Field(None, description="CV folder path override")
    append: bool = Field(False, description="Append to existing Excel instead of overwriting")


class ShortlistCriteria(BaseModel):
    """Criteria for shortlisting candidates."""
    min_experience: Optional[float] = Field(None, description="Minimum years of experience")
    max_experience: Optional[float] = Field(None, description="Maximum years of experience")
    min_age: Optional[int] = Field(None, description="Minimum age")
    max_age: Optional[int] = Field(None, description="Maximum age")
    required_skills: Optional[list[str]] = Field(None, description="Required skills (any match)")
    required_education: Optional[list[str]] = Field(None, description="Education keywords")
    required_job_title: Optional[list[str]] = Field(None, description="Job title keywords")
    required_location: Optional[list[str]] = Field(None, description="Location keywords")


class JobMatchRequest(BaseModel):
    """Request body for job matching."""
    job_description: str = Field(..., description="Full job description text")
    provider: Optional[str] = Field(None, description="LLM provider override")
    use_shortlisted: bool = Field(True, description="Use shortlisted candidates (False = use all)")


# ──────────────────────────────────────────────
# Email endpoints
# ──────────────────────────────────────────────

@router.post("/email/download", tags=["Email"])
async def download_emails(request: EmailDownloadRequest = EmailDownloadRequest()):
    """Trigger email download — fetches new CV attachments from the configured mailbox."""
    from src.email_downloader import download_attachments

    try:
        result = download_attachments(
            imap_server=request.imap_server,
            imap_port=request.imap_port,
            email_address=request.email_address,
            email_password=request.email_password,
            email_folder=request.email_folder,
        )
        return {"status": "success", "data": result}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error("Email download failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/email/status", tags=["Email"])
async def email_status():
    """Get the current email tracker status (last run, total downloaded, etc.)."""
    from src.email_downloader import get_tracker_status

    return {"status": "success", "data": get_tracker_status()}


@router.post("/email/reset-tracker", tags=["Email"])
async def reset_email_tracker():
    """Reset the email tracker to re-process all emails."""
    from src.email_downloader import reset_tracker

    reset_tracker()
    return {"status": "success", "message": "Email tracker has been reset."}


# ──────────────────────────────────────────────
# Extract endpoints
# ──────────────────────────────────────────────

@router.post("/extract", tags=["Extraction"])
async def extract_cvs(request: ExtractRequest = ExtractRequest()):
    """Run CV data extraction on all files in the CV folder."""
    from src.pipeline import run_extract

    try:
        folder = Path(request.folder) if request.folder else None
        output_file = run_extract(
            cv_folder=folder,
            provider=request.provider,
            append=request.append,
        )
        return {
            "status": "success",
            "output_file": str(output_file),
            "message": f"Extraction complete. Results saved to {output_file}",
        }
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        logger.error("Extraction failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/extract/files", tags=["Extraction"])
async def list_cv_files():
    """List all CV files currently in the CV folder."""
    from src.pipeline import scan_cv_folder

    try:
        files = scan_cv_folder(CV_FOLDER)
        return {
            "status": "success",
            "folder": str(CV_FOLDER),
            "total_files": len(files),
            "files": [
                {
                    "name": f.name,
                    "extension": f.suffix,
                    "size_bytes": f.stat().st_size,
                }
                for f in files
            ],
        }
    except FileNotFoundError:
        return {"status": "success", "folder": str(CV_FOLDER), "total_files": 0, "files": []}


# ──────────────────────────────────────────────
# Candidates endpoints
# ──────────────────────────────────────────────

@router.get("/candidates", tags=["Candidates"])
async def get_all_candidates():
    """Get all extracted candidates from the main Excel file."""
    from src.excel_writer import read_from_excel

    records = read_from_excel(OUTPUT_EXCEL)
    return {
        "status": "success",
        "total": len(records),
        "candidates": records,
    }


@router.post("/candidates/shortlist", tags=["Candidates"])
async def shortlist_candidates(criteria: ShortlistCriteria):
    """Apply shortlisting criteria to filter candidates."""
    from src.shortlister import shortlist_candidates as do_shortlist

    criteria_dict = criteria.model_dump(exclude_none=True)

    if not criteria_dict:
        raise HTTPException(status_code=400, detail="At least one criterion must be provided.")

    try:
        result = do_shortlist(criteria=criteria_dict)
        # Remove heavy candidate data from summary, keep just the key fields
        summary_candidates = []
        for c in result.get("candidates", []):
            summary_candidates.append({
                "full_name": c.get("full_name"),
                "email": c.get("email"),
                "phone": c.get("phone"),
                "current_job_title": c.get("current_job_title"),
                "experience_years": c.get("experience_years"),
                "age": c.get("age"),
                "skills": c.get("skills"),
                "location": c.get("location"),
                "_filename": c.get("_filename"),
            })

        return {
            "status": "success",
            "total": result["total"],
            "shortlisted": result["shortlisted"],
            "rejected": result["rejected"],
            "criteria_used": result["criteria_used"],
            "candidates": summary_candidates,
            "output_file": result.get("output_file"),
        }
    except Exception as e:
        logger.error("Shortlisting failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/candidates/shortlisted", tags=["Candidates"])
async def get_shortlisted_candidates():
    """Get candidates from the shortlisted Excel file."""
    from src.excel_writer import read_from_excel

    if not SHORTLIST_EXCEL.exists():
        return {"status": "success", "total": 0, "candidates": [], "message": "No shortlist found. Run shortlisting first."}

    records = read_from_excel(SHORTLIST_EXCEL)
    return {
        "status": "success",
        "total": len(records),
        "candidates": records,
    }


# ──────────────────────────────────────────────
# Job Match endpoints
# ──────────────────────────────────────────────

@router.post("/job-match", tags=["Job Matching"])
async def run_job_match(request: JobMatchRequest):
    """Run LLM-based job fit analysis on candidates."""
    from src.job_matcher import match_candidates_to_job
    from src.excel_writer import read_from_excel

    # Load candidates
    if request.use_shortlisted and SHORTLIST_EXCEL.exists():
        candidates = read_from_excel(SHORTLIST_EXCEL)
        source = "shortlisted"
    else:
        candidates = read_from_excel(OUTPUT_EXCEL)
        source = "all"

    if not candidates:
        raise HTTPException(status_code=404, detail="No candidates found. Run extraction first.")

    try:
        result = match_candidates_to_job(
            job_description=request.job_description,
            candidates=candidates,
            provider=request.provider,
        )

        # Simplify results for API response
        api_results = []
        for r in result.get("results", []):
            api_results.append({
                "full_name": r["candidate"].get("full_name"),
                "filename": r["candidate"].get("_filename"),
                "overall_score": r["match"].get("overall_score"),
                "recommendation": r["match"].get("recommendation"),
                "match_summary": r["match"].get("match_summary"),
                "strengths": r["match"].get("strengths"),
                "weaknesses": r["match"].get("weaknesses"),
                "key_matching_skills": r["match"].get("key_matching_skills"),
                "missing_requirements": r["match"].get("missing_requirements"),
                "status": r["status"],
            })

        return {
            "status": "success",
            "source": source,
            "total_candidates": result["total_candidates"],
            "avg_score": result["avg_score"],
            "strong_matches": result["strong_matches"],
            "good_matches": result["good_matches"],
            "partial_matches": result["partial_matches"],
            "poor_matches": result["poor_matches"],
            "output_file": result.get("output_file"),
            "results": api_results,
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error("Job matching failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/job-match/results", tags=["Job Matching"])
async def get_job_match_results():
    """Get the latest job match results from Excel."""
    from src.excel_writer import read_from_excel

    if not JOB_MATCH_EXCEL.exists():
        return {"status": "success", "total": 0, "results": [], "message": "No results. Run job matching first."}

    # Read from the job match Excel (different format than CV data)
    from openpyxl import load_workbook

    wb = load_workbook(str(JOB_MATCH_EXCEL), read_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        record = {}
        for i, value in enumerate(row):
            if i < len(headers):
                record[headers[i]] = value if value is not None else ""
        results.append(record)

    wb.close()
    return {"status": "success", "total": len(results), "results": results}


# ──────────────────────────────────────────────
# Config / Health endpoints
# ──────────────────────────────────────────────

@router.get("/health", tags=["System"])
async def health_check():
    """Health check endpoint."""
    return {"status": "healthy", "service": "HR Management System"}


@router.get("/config", tags=["System"])
async def get_config():
    """Get current runtime configuration (no secrets)."""
    return {
        "llm_provider": LLM_PROVIDER,
        "cv_folder": str(CV_FOLDER),
        "output_excel": str(OUTPUT_EXCEL),
        "shortlist_excel": str(SHORTLIST_EXCEL),
        "job_match_excel": str(JOB_MATCH_EXCEL),
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
    }
