"""
job_matcher.py — LLM-based job fit analysis.

Takes shortlisted candidates and a job description, sends each pair
to the LLM, and returns structured fit scores and analysis.

Output per candidate:
  - overall_score (0-100)
  - match_summary (brief text)
  - strengths (list)
  - weaknesses (list)
  - recommendation: "strong_match" / "good_match" / "partial_match" / "poor_match"
"""

import json
import logging
import re
import time
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from src.config import (
    LLM_PROVIDER,
    CLAUDE_API_KEY,
    CLAUDE_MODEL,
    GEMINI_API_KEY,
    GEMINI_MODEL,
    JOB_MATCH_EXCEL,
)

logger = logging.getLogger(__name__)


# ──────────────────────────────────────────────
# Prompt
# ──────────────────────────────────────────────

JOB_MATCH_SYSTEM_PROMPT = """You are an expert HR recruiter AI. Your job is to evaluate how well a candidate fits a given job description.

Analyze the candidate's profile against the job description and return a JSON object with exactly these fields:

{
  "overall_score": <integer 0-100>,
  "match_summary": "<2-3 sentence summary of the fit>",
  "strengths": ["<strength 1>", "<strength 2>", ...],
  "weaknesses": ["<weakness 1>", "<weakness 2>", ...],
  "recommendation": "<one of: strong_match, good_match, partial_match, poor_match>",
  "key_matching_skills": ["<skill 1>", "<skill 2>", ...],
  "missing_requirements": ["<requirement 1>", "<requirement 2>", ...]
}

Scoring guide:
- 80-100: Strong match — candidate meets most/all requirements
- 60-79:  Good match — candidate meets core requirements with minor gaps
- 40-59:  Partial match — candidate has relevant experience but significant gaps
- 0-39:   Poor match — candidate does not meet key requirements

Rules:
- Be objective and thorough in your assessment.
- Consider both explicit requirements and implied qualifications.
- Factor in experience level, skills, education, and domain relevance.
- Return ONLY the JSON object. No markdown, no explanation, no code blocks."""


def _build_match_prompt(candidate: dict[str, Any], job_description: str) -> str:
    """Build the user prompt for job matching."""
    candidate_info = "\n".join(
        f"  {key}: {value}"
        for key, value in candidate.items()
        if not key.startswith("_") and value and str(value).strip()
    )
    return f"""## Candidate Profile:
{candidate_info}

## Job Description:
{job_description}

Evaluate this candidate's fit for the job and return the JSON analysis."""


# ──────────────────────────────────────────────
# LLM Calls
# ──────────────────────────────────────────────

def _match_with_claude(candidate: dict, job_description: str) -> dict[str, Any]:
    """Score a candidate using Claude."""
    import anthropic

    if not CLAUDE_API_KEY or CLAUDE_API_KEY == "your_claude_api_key_here":
        raise ValueError("CLAUDE_API_KEY is not set.")

    client = anthropic.Anthropic(api_key=CLAUDE_API_KEY)
    response = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=2048,
        system=JOB_MATCH_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": _build_match_prompt(candidate, job_description)}],
    )
    return _parse_match_response(response.content[0].text)


def _match_with_gemini(candidate: dict, job_description: str) -> dict[str, Any]:
    """Score a candidate using Gemini."""
    import google.generativeai as genai

    if not GEMINI_API_KEY or GEMINI_API_KEY == "your_gemini_api_key_here":
        raise ValueError("GEMINI_API_KEY is not set.")

    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel(
        model_name=GEMINI_MODEL,
        system_instruction=JOB_MATCH_SYSTEM_PROMPT,
    )
    response = model.generate_content(_build_match_prompt(candidate, job_description))
    return _parse_match_response(response.text)


def _parse_match_response(raw: str) -> dict[str, Any]:
    """Parse the LLM's job match response into a dict."""
    cleaned = raw.strip()

    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*\n?", "", cleaned)
        cleaned = re.sub(r"\n?```\s*$", "", cleaned)

    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError as e:
        logger.error("Failed to parse job match response: %s", e)
        data = {
            "overall_score": 0,
            "match_summary": "Failed to parse LLM response",
            "strengths": [],
            "weaknesses": [],
            "recommendation": "error",
            "key_matching_skills": [],
            "missing_requirements": [],
            "_parse_error": str(e),
            "_raw_response": raw,
        }

    # Ensure all fields exist
    defaults = {
        "overall_score": 0,
        "match_summary": "",
        "strengths": [],
        "weaknesses": [],
        "recommendation": "unknown",
        "key_matching_skills": [],
        "missing_requirements": [],
    }
    for key, default in defaults.items():
        if key not in data:
            data[key] = default

    return data


# ──────────────────────────────────────────────
# Excel output for match results
# ──────────────────────────────────────────────

def _write_match_results_to_excel(
    results: list[dict[str, Any]],
    output_path: Path,
) -> Path:
    """Write job match results to a styled Excel file."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Job Match Results"

    # Headers
    headers = [
        "File Name", "Full Name", "Current Job Title", "Experience (Years)",
        "Overall Score", "Recommendation", "Match Summary",
        "Strengths", "Weaknesses", "Matching Skills", "Missing Requirements",
    ]

    # Styling
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Score-based fills
    score_fills = {
        "strong_match": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
        "good_match": PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid"),
        "partial_match": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
        "poor_match": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
    }

    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    ws.freeze_panes = "A2"

    # Column widths
    widths = [25, 22, 25, 12, 12, 16, 45, 40, 40, 35, 35]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w

    # Write data
    for row_idx, result in enumerate(results, 2):
        candidate = result.get("candidate", {})
        match_data = result.get("match", {})

        values = [
            candidate.get("_filename", ""),
            candidate.get("full_name", ""),
            candidate.get("current_job_title", ""),
            candidate.get("experience_years", ""),
            match_data.get("overall_score", 0),
            match_data.get("recommendation", ""),
            match_data.get("match_summary", ""),
            ", ".join(match_data.get("strengths", [])),
            ", ".join(match_data.get("weaknesses", [])),
            ", ".join(match_data.get("key_matching_skills", [])),
            ", ".join(match_data.get("missing_requirements", [])),
        ]

        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = cell_align
            cell.border = border

        # Color-code the recommendation cell
        rec = match_data.get("recommendation", "")
        if rec in score_fills:
            ws.cell(row=row_idx, column=6).fill = score_fills[rec]

    ws.auto_filter.ref = ws.dimensions
    wb.save(str(output_path))
    logger.info("Job match results saved: %s", output_path)
    return output_path


# ──────────────────────────────────────────────
# Public API
# ──────────────────────────────────────────────

def match_candidates_to_job(
    job_description: str,
    candidates: list[dict[str, Any]],
    provider: str | None = None,
    output_path: Path | None = None,
) -> dict[str, Any]:
    """
    Evaluate each candidate against a job description using the LLM.

    Args:
        job_description: The full job description text.
        candidates: List of candidate dicts (from shortlister or extraction).
        provider: LLM provider override ("claude" or "gemini").
        output_path: Path to save results Excel (default: JOB_MATCH_EXCEL).

    Returns:
        Summary dict with all results.
    """
    active_provider = (provider or LLM_PROVIDER).lower()
    out = output_path or JOB_MATCH_EXCEL

    if not job_description or not job_description.strip():
        raise ValueError("Job description cannot be empty.")

    if not candidates:
        logger.warning("No candidates to match.")
        return {"total": 0, "results": []}

    logger.info("Matching %d candidates against job description using %s...", len(candidates), active_provider)

    results = []
    for i, candidate in enumerate(candidates):
        name = candidate.get("full_name", f"Candidate {i + 1}")
        logger.info("  [%d/%d] Evaluating: %s", i + 1, len(candidates), name)

        try:
            if active_provider == "claude":
                match_data = _match_with_claude(candidate, job_description)
            elif active_provider == "gemini":
                match_data = _match_with_gemini(candidate, job_description)
            else:
                raise ValueError(f"Unknown provider: {active_provider}")

            results.append({
                "candidate": candidate,
                "match": match_data,
                "status": "success",
            })

            score = match_data.get("overall_score", 0)
            rec = match_data.get("recommendation", "unknown")
            logger.info("    → Score: %d/100 | %s", score, rec)

        except Exception as e:
            logger.error("    ✗ Error matching %s: %s", name, e)
            results.append({
                "candidate": candidate,
                "match": {
                    "overall_score": 0,
                    "match_summary": f"Error: {e}",
                    "strengths": [],
                    "weaknesses": [],
                    "recommendation": "error",
                    "key_matching_skills": [],
                    "missing_requirements": [],
                },
                "status": f"error: {e}",
            })

        # Small delay to respect rate limits
        if i < len(candidates) - 1:
            time.sleep(1)

    # Sort results by score (highest first)
    results.sort(key=lambda r: r["match"].get("overall_score", 0), reverse=True)

    # Write to Excel
    _write_match_results_to_excel(results, out)

    # Summary stats
    scores = [r["match"].get("overall_score", 0) for r in results if r["status"] == "success"]
    summary = {
        "total_candidates": len(candidates),
        "evaluated": len([r for r in results if r["status"] == "success"]),
        "errors": len([r for r in results if r["status"] != "success"]),
        "avg_score": round(sum(scores) / len(scores), 1) if scores else 0,
        "strong_matches": len([r for r in results if r["match"].get("recommendation") == "strong_match"]),
        "good_matches": len([r for r in results if r["match"].get("recommendation") == "good_match"]),
        "partial_matches": len([r for r in results if r["match"].get("recommendation") == "partial_match"]),
        "poor_matches": len([r for r in results if r["match"].get("recommendation") == "poor_match"]),
        "output_file": str(out),
        "results": results,
    }

    logger.info("=" * 50)
    logger.info("Job Match Summary:")
    logger.info("  Evaluated      : %d / %d", summary["evaluated"], summary["total_candidates"])
    logger.info("  Avg Score      : %.1f / 100", summary["avg_score"])
    logger.info("  Strong matches : %d", summary["strong_matches"])
    logger.info("  Good matches   : %d", summary["good_matches"])
    logger.info("  Partial matches: %d", summary["partial_matches"])
    logger.info("  Poor matches   : %d", summary["poor_matches"])
    logger.info("=" * 50)

    return summary
